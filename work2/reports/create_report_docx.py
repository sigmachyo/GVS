from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).parent
PROJECT_ROOT = ROOT.parent
OUTPUT = ROOT / "temp" / "report.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)


def set_run_font(run, name="Times New Roman", size=14, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    set_run_font(run)


def add_heading(document, text, level=1, centered=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size=14, bold=True)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Mm(12.5)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_code(document, title, path):
    add_heading(document, title, level=2)
    source = (PROJECT_ROOT / path).read_text(encoding="utf-8")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Mm(5)
    paragraph.paragraph_format.right_indent = Mm(5)
    for line in source.splitlines():
        run = paragraph.add_run(line + "\n")
        set_run_font(run, name="Courier New", size=8)


def add_figure_or_placeholder(document, filename, caption):
    image_path = ROOT / "temp" / filename
    if image_path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Mm(165))
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("[График будет добавлен после запуска бенчмарка на CUDA-машине]")
        set_run_font(run, size=12)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(caption)
    set_run_font(run)


document = Document()
section = document.sections[0]
section.top_margin = Mm(20)
section.bottom_margin = Mm(20)
section.left_margin = Mm(30)
section.right_margin = Mm(10)
section.header_distance = Mm(10)
section.footer_distance = Mm(10)
add_page_number(section.footer.paragraphs[0])

styles = document.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(14)

# Титульный лист по приложению М СТУ 7.5-07-2021.
for text, bold, size in [
    ("Министерство науки и высшего образования Российской Федерации", False, 14),
    ("Федеральное государственное автономное образовательное учреждение", False, 14),
    ("высшего образования", False, 14),
    ("«СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»", True, 14),
    ("", False, 14),
    ("Институт: [указать полное наименование института]", False, 14),
    ("Кафедра: [указать полное наименование кафедры]", False, 14),
    ("", False, 14),
    ("ОТЧЕТ О ПРАКТИЧЕСКОЙ РАБОТЕ", True, 16),
    ("Тема: «Умножение матриц на CUDA с использованием паттерна Data + View»", True, 14),
    ("", False, 14),
    ("Преподаватель: [инициалы, фамилия]", False, 14),
    ("Студент: [номер группы, инициалы, фамилия]", False, 14),
    ("", False, 14),
    ("Красноярск", False, 14),
    ("2026", False, 14),
]:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if not text.startswith(("Преподаватель", "Студент")) else WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
document.add_page_break()

add_heading(document, "ВВЕДЕНИЕ", centered=True)
add_body(document, "Цель работы -- освоить двумерные сетки нитей CUDA и реализовать наивное умножение матриц C = A * B с использованием паттерна Data + View.")
add_body(document, "Для достижения цели необходимо разработать шаблонные классы Data, MatrixView и Matrix, реализовать CUDA-кернел наивного умножения матриц, перегрузить operator*, проверить результаты с помощью Google Test и Eigen, а также исследовать производительность с помощью Google Benchmark.")

add_heading(document, "1 Разработка программного решения")
add_heading(document, "1.1 Класс Data", level=2)
add_body(document, "Класс Data владеет массивом в глобальной памяти устройства. Выделение памяти выполняется в конструкторе, освобождение -- в деструкторе, что соответствует идиоме RAII. Для матрицы память выделяется одним линейным блоком размером nrows * ncols элементов.")
add_code(document, "Реализация Data", "core/include/Data.cuh")
add_heading(document, "1.2 Класс MatrixView", level=2)
add_body(document, "Класс MatrixView хранит указатель на данные, количество строк и столбцов. Элементы хранятся в порядке строк (row-major): элемент [i][j] расположен по адресу data_[i * ncols_ + j]. Объекты передаются в CUDA-кернел по значению.")
add_code(document, "Реализация MatrixView", "core/include/MatrixView.cuh")
add_heading(document, "1.3 Класс Matrix", level=2)
add_body(document, "Класс Matrix является фасадом: агрегирует Data через std::shared_ptr и содержит собственное представление MatrixView. Метод matmul_from запускает CUDA-кернел с двумерной сеткой нитей.")
add_code(document, "Реализация Matrix", "core/include/Matrix.cuh")
add_heading(document, "1.4 CUDA-кернел умножения", level=2)
add_body(document, "Кернел принимает три объекта MatrixView по значению. Координаты row и col вычисляются через blockIdx и threadIdx. Каждый поток вычисляет один элемент результирующей матрицы, накапливая скалярное произведение строки A и столбца B.")
add_code(document, "Реализация kernel_matmul_naive", "core/include/kernel_matmul_naive.cuh")

add_heading(document, "2 Тестирование корректности")
add_body(document, "Для проверки operator* используются 343 комбинации размеров (m, k, n) из множества {1, 2, 3, 127, 128, 129, 512}. Эталонный результат вычисляется с помощью Eigen с row-major маппингом. Проверка выполняется через isApprox с точностью 10^-5.")

add_heading(document, "3 Экспериментальное исследование производительности")
add_body(document, "Бенчмарки выполняются для квадратных матриц n x n, где n ∈ {16, 32, 64, 128, 256, 512, 1024}. Для CUDA используются события CUDA Events API. Память выделяется и заполняется до начала измеряемого цикла.")
add_heading(document, "3.1 График реальной вычислительной сложности", level=2)
add_figure_or_placeholder(document, "real_complexity.png", "Рисунок 1 -- Реальная вычислительная сложность умножения матриц")
add_heading(document, "3.2 График ускорения", level=2)
add_figure_or_placeholder(document, "speedup.png", "Рисунок 2 -- Ускорение CUDA относительно Eigen")

add_heading(document, "4 Интерпретация результатов")
add_body(document, "Теоретическая сложность наивного умножения матриц кубична O(n³): для каждого из n² элементов результата выполняется n операций умножения и сложения. На малых размерах CUDA уступает Eigen из-за постоянных накладных расходов запуска кернела и синхронизации. На больших размерах CUDA показывает значительное ускорение за счёт параллельного выполнения.")

add_heading(document, "ЗАКЛЮЧЕНИЕ", centered=True)
add_body(document, "В работе разработана шаблонная реализация умножения матриц на CUDA с использованием паттерна Data + View. Класс Data управляет памятью устройства по RAII, MatrixView предоставляет тривиально копируемое двумерное представление данных в row-major порядке, а Matrix объединяет эти компоненты и предоставляет перегруженный operator*.")
add_body(document, "Корректность решения проверяется сравнением с Eigen на 343 комбинациях размеров. Методика бенчмаркинга использует CUDA Events API и исключает выделение памяти из измеряемого участка.")

add_heading(document, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", centered=True)
sources = [
    "Стандарт университета. Система менеджмента качества. Общие требования к построению, изложению и оформлению документов учебной деятельности : СТУ 7.5-07-2021. -- Красноярск, 2021.",
    "CUDA C++ Programming Guide [Электронный ресурс]. -- URL: https://docs.nvidia.com/cuda/cuda-c-programming-guide/ (дата обращения: 11.10.2026).",
    "Eigen documentation [Электронный ресурс]. -- URL: https://eigen.tuxfamily.org/dox/ (дата обращения: 11.10.2026).",
    "GoogleTest documentation [Электронный ресурс]. -- URL: https://google.github.io/googletest/ (дата обращения: 11.10.2026).",
    "Google Benchmark documentation [Электронный ресурс]. -- URL: https://github.com/google/benchmark (дата обращения: 11.10.2026).",
]
for index, source in enumerate(sources, 1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Mm(12.5)
    run = paragraph.add_run(f"{index}. {source}")
    set_run_font(run)

document.save(OUTPUT)
print(OUTPUT)
